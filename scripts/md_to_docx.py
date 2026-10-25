"""把使用教程 Markdown 转成排版好的 DOCX（供再导出 PDF）。"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

BODY_FONT = "Microsoft YaHei"
MONO_FONT = "SimSun"  # 宋体自带方框线与等宽中文，ASCII 示意图才对得齐
ACCENT = RGBColor(0x3E, 0x6C, 0x96)
MUTED = RGBColor(0x7D, 0x76, 0x6C)


def set_run_font(run, font_name=BODY_FONT, size=None, bold=None, color=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(paragraph, text, size=10.5, color=None):
    """处理 **粗体** 与 `代码`。"""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, font_name=MONO_FONT, size=size - 0.5, color=RGBColor(0x8B, 0x4A, 0x1F))
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color)


def shade(cell_or_paragraph, hex_color):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), hex_color)
    cell_or_paragraph.append(el)


def build(md_path: Path, out_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.line_spacing = 1.38
    normal.paragraph_format.space_after = Pt(3)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # 代码块
        if line.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            shade(p._p.get_or_add_pPr(), "F2ECE2")
            for idx, code_line in enumerate(block):
                if idx:
                    p.add_run().add_break()
                run = p.add_run(code_line.replace("\t", "    "))
                set_run_font(run, font_name=MONO_FONT, size=9)
            continue

        # 表格
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            header = [c.strip() for c in line.strip().strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1

            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for idx, text in enumerate(header):
                cell = table.rows[0].cells[idx]
                cell.text = ""
                para = cell.paragraphs[0]
                run = para.add_run(text)
                set_run_font(run, size=9.5, bold=True)
                shade(cell._tc.get_or_add_tcPr(), "F2ECE2")
            for row in rows:
                cells = table.add_row().cells
                for idx, text in enumerate(row[: len(header)]):
                    cells[idx].text = ""
                    add_inline(cells[idx].paragraphs[0], text, size=9.5)
            doc.add_paragraph()
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1:
                h = doc.add_heading(level=0)
                run = h.add_run(text)
                set_run_font(run, size=20, bold=True, color=ACCENT)
                h.paragraph_format.space_after = Pt(10)
            else:
                h = doc.add_heading(level=level - 1)
                run = h.add_run(text)
                size = {2: 14, 3: 12, 4: 11}[level]
                set_run_font(run, size=size, bold=True, color=ACCENT if level == 2 else None)
                h.paragraph_format.space_before = Pt(12 if level == 2 else 8)
                h.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # 引用
        if line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            add_inline(p, line.lstrip("> ").strip(), size=10, color=MUTED)
            i += 1
            continue

        # 分隔线
        if line.strip() == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run("─" * 40)
            set_run_font(run, size=8, color=MUTED)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # 列表
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            indent = len(m.group(1)) // 2
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.75 + indent * 0.6)
            p.paragraph_format.line_spacing = 1.4
            add_inline(p, m.group(2))
            i += 1
            continue

        # 有序列表：手动写编号，避免 python-docx 的 List Number 跨节连续计数
        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if m:
            counter = 0
            while i < len(lines):
                mm = re.match(r"^(\s*)\d+\.\s+(.*)$", lines[i])
                if not mm:
                    break
                counter += 1
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.95)
                p.paragraph_format.first_line_indent = Cm(-0.45)
                p.paragraph_format.line_spacing = 1.4
                p.paragraph_format.space_after = Pt(3)
                run = p.add_run(f"{counter}. ")
                set_run_font(run, size=10.5, bold=True)
                add_inline(p, mm.group(2))
                i += 1
            continue

        # 空行
        if not line.strip():
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_inline(p, line.strip())
        i += 1

    doc.save(out_path)
    print(f"saved: {out_path}")


if __name__ == "__main__":
    build(Path(sys.argv[1]), Path(sys.argv[2]))
